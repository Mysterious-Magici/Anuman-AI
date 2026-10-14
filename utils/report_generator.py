"""Report rendering: PDF (ReportLab, logo embedded), DOCX, Markdown and plain text."""
from __future__ import annotations

import io
import re
from xml.sax.saxutils import escape

import pandas as pd

# pyrefly: ignore [missing-import]
from config import APP_NAME
# pyrefly: ignore [missing-import]
from models import AnalysisResult, ReportData

PRIMARY_HEX = "#4F46E5"
PDF_MAX_ROWS, DOC_MAX_ROWS, PDF_MAX_COLS = 200, 500, 10


class ReportError(Exception):
    pass


METHODOLOGY = {
    "summarize": "Text was extracted locally, split into chunks and summarized by Google Gemini (long inputs use a map-reduce pass).",
    "key_points": "Text was extracted locally, split into chunks and key points were drawn by Google Gemini.",
    "qa": "Relevant passages were retrieved locally by keyword scoring and sent to Google Gemini with a request to answer only from them.",
    "extract": "Relevant passages were selected locally; Google Gemini returned JSON that was validated and converted to a table.",
    "compare": "Representative excerpts of each document were selected locally and compared by Google Gemini.",
    "data_profile": "Counts and statistics were computed locally with pandas. Any AI commentary was produced by Google Gemini from the profile and a small sample.",
}


# ---------- building report data ----------

def default_summary(content: str, limit: int = 600) -> str:
    for line in content.splitlines():
        s = line.strip()
        if len(s) > 40 and not s.startswith(("#", "|", "-", "*", "[")):
            return re.sub(r"\*\*|`", "", s)[:limit]
    return re.sub(r"[#*`|]", "", content).strip()[:limit]


def split_sections(md: str) -> list[tuple[str, str]]:
    sections, heading, buf = [], "Overview", []
    for line in md.splitlines():
        m = re.match(r"^#{1,3}\s+(.*)$", line)
        if m:
            if "\n".join(buf).strip():
                sections.append((heading, "\n".join(buf).strip()))
            heading, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if "\n".join(buf).strip():
        sections.append((heading, "\n".join(buf).strip()))
    return sections or [("Findings", md.strip())]


def build_report_data(result: AnalysisResult, doc_infos: list[dict], title: str | None = None,
                      summary: str | None = None, include_table: bool = True,
                      include_sources: bool = True, include_methodology: bool = True) -> ReportData:
    method = ""
    if include_methodology:
        lines = [f"- {METHODOLOGY.get(result.task, 'Analysis was performed with the configured AI provider.')}"]
        if result.providers:
            lines.append(f"- Providers used: {', '.join(result.providers)}.")
        if result.coverage is not None:
            lines.append(f"- Share of the selected text sent to the model: {result.coverage:.0%}.")
        lines += [f"- Limitation: {w}" for w in result.warnings]
        lines.append("- AI output can contain errors; verify important facts against the source documents.")
        method = "\n".join(lines)
    raw_title = title or (result.title if result else "") or "Report"
    return ReportData(
        title=raw_title.strip() or "Report",
        summary=(summary if summary is not None else default_summary(result.content)).strip(),
        sections=split_sections(result.content),
        documents=doc_infos,
        table=result.table if include_table else None,
        sources=result.sources if include_sources else [],
        methodology=method,
    )


# ---------- shared structure ----------

_SEP = re.compile(r"^:?-{2,}:?$")


def parse_blocks(md: str) -> list[tuple]:
    blocks, lines, i = [], md.splitlines(), 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not (cells and all(_SEP.match(c) for c in cells)):
                    rows.append(cells)
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        if m := re.match(r"^#{1,6}\s+(.*)", line):
            blocks.append(("heading", m.group(1)))
        elif m := re.match(r"^\s*[-*\u2022]\s+(.*)", line):
            blocks.append(("bullet", m.group(1)))
        elif m := re.match(r"^\s*(\d+)[.)]\s+(.*)", line):
            blocks.append(("number", m.group(1), m.group(2)))
        else:
            blocks.append(("para", line.strip()))
        i += 1
    return blocks


def _has_table(report: ReportData) -> bool:
    return report.table is not None and not report.table.empty


def _df_rows(df: pd.DataFrame, max_rows: int, max_cols: int | None = None):
    note, d = "", df
    if max_cols and d.shape[1] > max_cols:
        note += f"Only the first {max_cols} of {d.shape[1]} columns are shown. "
        d = d.iloc[:, :max_cols]
    if len(d) > max_rows:
        note += f"Only the first {max_rows} of {len(df)} rows are shown. "
        d = d.head(max_rows)
    if note:
        note += "Download the CSV/XLSX export for the complete table."
    body = d.astype(object).where(d.notna(), "").astype(str).values.tolist()
    return [[str(c) for c in d.columns]] + body, note


def _outline(report: ReportData, max_rows: int, max_cols: int | None = None):
    sections, n = [], [1]

    def add(heading, items):
        sections.append((f"{n[0]}. {heading}", items))
        n[0] += 1

    add("Executive Summary", [("md", report.summary.strip() or "No summary provided.")])
    if report.documents:
        cols = list(dict.fromkeys(k for d in report.documents for k in d))
        add("Document Information", [("table", [cols] + [[str(d.get(c, "")) for c in cols] for d in report.documents])])
    for heading, body in report.sections:
        add(heading, [("md", body)])
    if _has_table(report):
        rows, note = _df_rows(report.table, max_rows, max_cols)
        add("Data Table", [("table", rows)] + ([("note", note)] if note else []))
    if report.sources:
        refs = report.sources[:100]
        md = "\n".join(f"- {r}" for r in refs)
        if len(report.sources) > 100:
            md += f"\n- ... and {len(report.sources) - 100} more"
        add("Source References", [("md", md)])
    if report.methodology.strip():
        add("Methodology and Limitations", [("md", report.methodology)])
    return sections


# ---------- Markdown / TXT ----------

def _rows_to_md(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    ncols = max(len(r) for r in rows)
    if ncols == 0:
        return ""
    def fmt(r):
        return "| " + " | ".join(str(c).replace("|", "\\|").replace("\n", " ") for c in r) + " |"
    rows = [r + [""] * (ncols - len(r)) for r in rows]
    return "\n".join([fmt(rows[0]), "| " + " | ".join(["---"] * ncols) + " |"] + [fmt(r) for r in rows[1:]])


def build_markdown(report: ReportData, max_rows: int = DOC_MAX_ROWS) -> str:
    lines = [f"# {report.title}", "", f"*{APP_NAME} - generated {report.generated_at:%Y-%m-%d %H:%M}*", ""]
    for heading, items in _outline(report, max_rows):
        lines += [f"## {heading}", ""]
        for kind, payload in items:
            if kind == "md":
                lines += [payload, ""]
            elif kind == "table":
                lines += [_rows_to_md(payload), ""]
            else:
                lines += [f"*{payload}*", ""]
    return "\n".join(lines)


def build_text(report: ReportData) -> str:
    out = []
    for line in build_markdown(report).splitlines():
        line = re.sub(r"^#{1,6}\s+", "", line)
        out.append(re.sub(r"\*\*|`", "", line))
    return "\n".join(out)


# ---------- DOCX ----------

_XML_BAD = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text) -> str:
    return _XML_BAD.sub("", str(text))


def _add_inline(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*)", _clean(text)):
        if len(part) > 4 and part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part.replace("`", ""))


def _add_page_field(paragraph) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    for kind in ("begin", "instr", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def build_docx(report: ReportData) -> bytes:
    try:
        from docx import Document
    except ImportError as exc:
        raise ReportError("python-docx is not installed. Run: pip install python-docx") from exc
    d = Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = _clean(f"{APP_NAME}  |  {report.title}"[:110])
    footer = section.footer.paragraphs[0]
    footer.add_run(f"Generated {report.generated_at:%Y-%m-%d}   |   Page ")
    _add_page_field(footer)
    d.add_heading(_clean(report.title), 0)
    d.add_paragraph(f"{APP_NAME} - generated {report.generated_at:%Y-%m-%d %H:%M}")
    for heading, items in _outline(report, DOC_MAX_ROWS):
        d.add_heading(_clean(heading), 1)
        for kind, payload in items:
            if kind == "md":
                for b in parse_blocks(payload):
                    if b[0] == "heading":
                        d.add_heading(_clean(b[1]), 2)
                    elif b[0] == "bullet":
                        _add_inline(d.add_paragraph(style="List Bullet"), b[1])
                    elif b[0] == "number":
                        _add_inline(d.add_paragraph(), f"{b[1]}. {b[2]}")
                    elif b[0] == "table":
                        _docx_table(d, b[1])
                    else:
                        _add_inline(d.add_paragraph(), b[1])
            elif kind == "table":
                _docx_table(d, payload)
            else:
                d.add_paragraph().add_run(_clean(payload)).italic = True
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _docx_table(d, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    if ncols == 0:
        return
    table = d.add_table(rows=1, cols=ncols)
    try:
        table.style = "Light Grid Accent 1"
    except Exception:
        pass
    for j in range(ncols):
        cell = table.rows[0].cells[j]
        cell.text = ""
        cell.paragraphs[0].add_run(_clean(rows[0][j] if j < len(rows[0]) else "")[:300]).bold = True
    for r in rows[1:]:
        cells = table.add_row().cells
        for j in range(ncols):
            cells[j].text = _clean(r[j] if j < len(r) else "")[:1000]
    d.add_paragraph()


# ---------- PDF ----------

def _pdf_safe(text) -> str:
    """Built-in PDF fonts cover Western characters only; others become '?'."""
    return str(text).encode("cp1252", errors="replace").decode("cp1252")


def _inline(text: str) -> str:
    t = escape(_pdf_safe(text))
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"<i>\1</i>", t)
    return re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", t)


def _logo(svg: str | None, height: float):
    """Convert the app's inline SVG logo into a scaled ReportLab drawing (vector, embedded in the PDF)."""
    if not svg:
        return None
    try:
        from svglib.svglib import svg2rlg
        drawing = svg2rlg(io.BytesIO(svg.encode("utf-8")))
        if drawing is None:
            return None
        scale = height / drawing.height
        drawing.width, drawing.height = drawing.width * scale, drawing.height * scale
        drawing.scale(scale, scale)
        return drawing
    except Exception:
        return None


def _bullet_paragraph(text: str, style, bullet_text: str):
    """Bulleted paragraph with a plain-text fallback if the inline markup is malformed."""
    from reportlab.platypus import Paragraph
    try:
        return Paragraph(_inline(text), style, bulletText=bullet_text)
    except Exception:
        plain = escape(_pdf_safe(re.sub(r"[*`]", "", text)))
        return Paragraph(plain, style, bulletText=bullet_text)


def build_pdf(report: ReportData, logo_svg: str | None) -> bytes:
    try:
        from reportlab.graphics import renderPDF
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ReportError("ReportLab is not installed. Run: pip install reportlab") from exc

    big, small = _logo(logo_svg, 18 * mm), _logo(logo_svg, 8 * mm)
    if logo_svg and (big is None or small is None):
        raise ReportError("The logo could not be embedded in the PDF. Check that 'svglib' is installed "
                          "(pip install svglib).")

    primary = colors.HexColor(PRIMARY_HEX)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=styles["BodyText"], fontSize=10, leading=14)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=14, bulletIndent=2)
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=15, textColor=primary, spaceBefore=14, spaceAfter=6)
    h2 = ParagraphStyle("H2", parent=body, fontName="Helvetica-Bold", fontSize=11, spaceBefore=6)
    title_style = ParagraphStyle("Title", parent=styles["Title"], fontSize=22, alignment=0, textColor=colors.HexColor("#111827"))
    note_style = ParagraphStyle("Note", parent=body, fontName="Helvetica-Oblique", fontSize=8, textColor=colors.HexColor("#6B7280"))
    page_w, page_h = A4
    avail = page_w - 36 * mm

    def para(markup: str, style):
        try:
            return Paragraph(markup, style)
        except Exception:  # malformed inline markup: fall back to plain text
            return Paragraph(escape(re.sub(r"<[^>]+>", "", markup)), style)

    def pdf_table(rows):
        if not rows:
            return Spacer(1, 1)
        ncols = max(len(r) for r in rows)
        if ncols == 0:
            return Spacer(1, 1)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        font = 8 if ncols <= 5 else 7 if ncols <= 8 else 6
        cell = ParagraphStyle("Cell", parent=body, fontSize=font, leading=font + 2)
        head = ParagraphStyle("Head", parent=cell, fontName="Helvetica-Bold", textColor=colors.white)
        data = [[para(_inline(str(c)[:300]), head if i == 0 else cell) for c in r] for i, r in enumerate(rows)]
        t = Table(data, colWidths=[avail / ncols] * ncols, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), primary),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D1D5DB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F9FAFB")]),
        ]))
        return t

    def decorate(canvas, doc):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#E5E7EB"))
        canvas.line(18 * mm, page_h - 19 * mm, page_w - 18 * mm, page_h - 19 * mm)
        x = 18 * mm
        if small is not None:
            renderPDF.draw(small, canvas, x, page_h - 17 * mm)
            x += small.width + 3 * mm
        canvas.setFillColor(primary)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(x, page_h - 13 * mm, _pdf_safe(APP_NAME))
        canvas.setFillColor(colors.HexColor("#6B7280"))
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(page_w - 18 * mm, page_h - 13 * mm, _pdf_safe(report.title)[:70])
        canvas.line(18 * mm, 14 * mm, page_w - 18 * mm, 14 * mm)
        canvas.drawString(18 * mm, 9 * mm, f"Generated {report.generated_at:%Y-%m-%d %H:%M}")
        canvas.drawRightString(page_w - 18 * mm, 9 * mm, f"Page {doc.page}")
        canvas.restoreState()

    story = []
    brand = para(f"<font size=16 color='{PRIMARY_HEX}'><b>{escape(_pdf_safe(APP_NAME))}</b></font><br/>"
                 "<font size=9 color='#6B7280'>Document Intelligence Report</font>", body)
    cover = Table([[big if big is not None else "", brand]], colWidths=[24 * mm, avail - 24 * mm])
    cover.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [cover, Spacer(1, 10), para(_inline(report.title), title_style),
              para(f"Generated {report.generated_at:%Y-%m-%d %H:%M} &nbsp;|&nbsp; "
                   f"Documents: {len(report.documents)}", note_style), Spacer(1, 6)]

    for heading, items in _outline(report, PDF_MAX_ROWS, PDF_MAX_COLS):
        story.append(para(_inline(heading), h1))
        for kind, payload in items:
            if kind == "table":
                story += [pdf_table(payload), Spacer(1, 6)]
            elif kind == "note":
                story.append(para(_inline(payload), note_style))
            else:
                for b in parse_blocks(payload):
                    if b[0] == "heading":
                        story.append(para(f"<b>{_inline(b[1])}</b>", h2))
                    elif b[0] == "bullet":
                        story.append(_bullet_paragraph(b[1], bullet, "\u2022"))
                    elif b[0] == "number":
                        story.append(_bullet_paragraph(b[2], bullet, f"{b[1]}."))
                    elif b[0] == "table":
                        story += [pdf_table(b[1]), Spacer(1, 6)]
                    else:
                        story.append(para(_inline(b[1]), body))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=26 * mm,
                      bottomMargin=20 * mm, title=_pdf_safe(report.title), author=APP_NAME
                      ).build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()