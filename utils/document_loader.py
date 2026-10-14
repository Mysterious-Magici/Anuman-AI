"""Local document ingestion (no API calls). Each loader returns located text segments and tables."""
from __future__ import annotations

import io
import os

import pandas as pd

# pyrefly: ignore [missing-import]
from models import LoadedDocument, Segment


class DocumentError(Exception):
    """A readable problem with a specific file."""


def ocr_available() -> bool:
    try:
        # pyrefly: ignore [missing-import]
        import pytesseract  # noqa: F401
        # pyrefly: ignore [missing-import]
        from pdf2image import convert_from_bytes  # noqa: F401
        return True
    except ImportError:
        return False


# ---------- helpers ----------

def _unique_headers(headers) -> list[str]:
    seen, out = {}, []
    for i, h in enumerate(headers):
        h = (str(h) if h is not None else "").strip() or f"col_{i + 1}"
        if h in seen:
            seen[h] += 1
            h = f"{h}_{seen[h]}"
        else:
            seen[h] = 1
        out.append(h)
    return out


def _rows_to_df(rows: list[list[str]]) -> pd.DataFrame:
    if not rows:
        return pd.DataFrame()
    max_cols = max(len(r) for r in rows)
    if max_cols == 0:
        return pd.DataFrame()
    padded_header = rows[0] + [f"col_{j+1}" for j in range(len(rows[0]), max_cols)]
    headers = _unique_headers(padded_header)
    if len(rows) == 1:
        return pd.DataFrame([], columns=headers)
    padded_body = [r + [""] * (max_cols - len(r)) for r in rows[1:]]
    return pd.DataFrame(padded_body, columns=headers)


def _df_blocks(df: pd.DataFrame, max_chars: int = 1500) -> list[tuple[int, int, str]]:
    """Split a table into CSV text blocks (header repeated) so no rows are dropped."""
    lines = df.to_csv(index=False).splitlines()
    if len(lines) < 2:
        return []
    header, body = lines[0], lines[1:]
    blocks, cur, size, start = [], [], len(header), 1
    for i, line in enumerate(body, start=1):
        if cur and size + len(line) + 1 > max_chars:
            blocks.append((start, i - 1, header + "\n" + "\n".join(cur)))
            cur, size, start = [], len(header), i
        cur.append(line)
        size += len(line) + 1
    if cur:
        blocks.append((start, len(body), header + "\n" + "\n".join(cur)))
    return blocks


def _table_document(name, ext, size, sheets: dict[str, pd.DataFrame], show_sheet: bool) -> LoadedDocument:
    doc = LoadedDocument(name, ext, size)
    total_rows = 0
    for sheet, df in sheets.items():
        if df.empty:
            doc.warnings.append(f"'{sheet}' has no data rows.")
            continue
        doc.tables[str(sheet)] = df
        total_rows += len(df)
        for a, b, text in _df_blocks(df):
            loc = f"sheet '{sheet}', rows {a}-{b}" if show_sheet else f"rows {a}-{b}"
            doc.segments.append(Segment(text, loc))
    first = next(iter(doc.tables.values()), pd.DataFrame())
    doc.metadata = {"tables": len(doc.tables), "rows": total_rows, "columns": first.shape[1]}
    return doc


def _decode(data: bytes) -> tuple[str, bool]:
    for enc in ("utf-8-sig", "utf-8"):
        try:
            return data.decode(enc), True
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1"), False


# ---------- loaders ----------

def _ocr_pdf(data: bytes):
    """Optional OCR. Returns None if OCR is unavailable or fails."""
    try:
        # pyrefly: ignore [missing-import]
        import pytesseract
        # pyrefly: ignore [missing-import]
        from pdf2image import convert_from_bytes
        out = []
        for i, img in enumerate(convert_from_bytes(data, dpi=200), start=1):
            text = pytesseract.image_to_string(img).strip()
            if text:
                out.append(Segment(text, f"page {i} (OCR)"))
        return out
    except Exception:
        return None


def _load_pdf(name: str, data: bytes) -> LoadedDocument:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise DocumentError("The PDF is password-protected.")
    doc = LoadedDocument(name, ".pdf", len(data), metadata={"pages": len(reader.pages)})
    empty = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            doc.segments.append(Segment(text, f"page {i}"))
        else:
            empty.append(i)
    if not doc.segments:
        ocr = _ocr_pdf(data)
        if not ocr:
            raise DocumentError("No extractable text found (the PDF is probably scanned images). "
                                "Install the optional OCR tools described in the README to read it.")
        doc.segments = ocr
        doc.warnings.append("Text was recognized with OCR and may contain errors.")
    elif empty:
        shown = ", ".join(map(str, empty[:15])) + ("..." if len(empty) > 15 else "")
        doc.warnings.append(f"No extractable text on page(s): {shown} (images or scans are not read).")
    return doc


def _load_docx(name: str, data: bytes) -> LoadedDocument:
    import docx
    d = docx.Document(io.BytesIO(data))
    doc = LoadedDocument(name, ".docx", len(data))
    heading, buf = "document start", []

    def flush():
        text = "\n".join(buf).strip()
        if text:
            doc.segments.append(Segment(text, f"section: {heading}"))
        buf.clear()

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style_name = p.style.name if p.style else ""
        if (style_name or "").lower().startswith(("heading", "title")):
            flush()
            heading = t[:80]
        buf.append(t)
    flush()
    for i, tbl in enumerate(d.tables, start=1):
        df = _rows_to_df([[c.text.strip() for c in r.cells] for r in tbl.rows])
        if not df.empty:
            doc.tables[f"table {i}"] = df
            doc.segments.append(Segment(df.to_csv(index=False), f"table {i}"))
    doc.metadata = {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}
    return doc


def _load_text(name: str, data: bytes) -> LoadedDocument:
    text, clean = _decode(data)
    text = text.strip()
    if not text:
        raise DocumentError("The file contains no text.")
    doc = LoadedDocument(name, os.path.splitext(name)[1].lower(), len(data),
                         segments=[Segment(text, "full text")],
                         metadata={"characters": len(text), "lines": text.count("\n") + 1})
    if not clean:
        doc.warnings.append("The file was not UTF-8; it was read as Latin-1 and some characters may look wrong.")
    return doc


def _load_csv(name: str, data: bytes) -> LoadedDocument:
    try:
        try:
            df = pd.read_csv(io.BytesIO(data))
        except UnicodeDecodeError:
            df = pd.read_csv(io.BytesIO(data), encoding="latin-1")
    except pd.errors.EmptyDataError as exc:
        raise DocumentError("The CSV file has no data.") from exc
    return _table_document(name, ".csv", len(data), {"data": df}, show_sheet=False)


def _load_xlsx(name: str, data: bytes) -> LoadedDocument:
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    doc = _table_document(name, ".xlsx", len(data), sheets, show_sheet=True)
    doc.metadata["sheets"] = list(sheets)
    return doc


def _load_pptx(name: str, data: bytes) -> LoadedDocument:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    doc = LoadedDocument(name, ".pptx", len(data), metadata={"slides": len(prs.slides)})
    for i, slide in enumerate(prs.slides, start=1):
        parts, k = [], 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                k += 1
                df = _rows_to_df([[c.text.strip() for c in r.cells] for r in shape.table.rows])
                doc.tables[f"slide {i} table {k}"] = df
                parts.append(df.to_csv(index=False))
        if getattr(slide, "has_notes_slide", False) and getattr(slide, "notes_slide", None) and getattr(slide.notes_slide, "notes_text_frame", None):
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Speaker notes: " + notes)
        if parts:
            doc.segments.append(Segment("\n".join(parts), f"slide {i}"))
    return doc


_LOADERS = {".pdf": _load_pdf, ".docx": _load_docx, ".txt": _load_text, ".md": _load_text,
            ".csv": _load_csv, ".xlsx": _load_xlsx, ".pptx": _load_pptx}


def load_document(name: str, data: bytes) -> LoadedDocument:
    if data is None:
        raise DocumentError("No file data provided.")
    ext = os.path.splitext(name)[1].lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise DocumentError(f"Unsupported file type '{ext}'.")
    try:
        doc = loader(name, data)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError(f"Could not read '{name}' ({type(exc).__name__}). The file may be corrupted.") from exc
    if not doc.segments and not doc.tables:
        raise DocumentError(f"No readable content was found in '{name}'.")
    return doc