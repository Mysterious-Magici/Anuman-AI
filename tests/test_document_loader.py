import io

import pandas as pd
import pytest

# pyrefly: ignore [missing-import]
from utils.document_loader import DocumentError, load_document


def test_txt_is_loaded_with_location():
    doc = load_document("n.txt", b"Hello world")
    assert doc.segments[0].text == "Hello world"
    assert doc.segments[0].location == "full text"


def test_empty_txt_raises():
    with pytest.raises(DocumentError):
        load_document("n.txt", b"   \n ")


def test_csv_keeps_table_and_row_ranges():
    doc = load_document("d.csv", b"a,b\n1,2\n3,4\n")
    assert list(doc.tables["data"].columns) == ["a", "b"]
    assert doc.segments[0].location == "rows 1-2"


def test_xlsx_sheets_become_tables():
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"x": [1, 2]}).to_excel(w, sheet_name="Sales", index=False)
    doc = load_document("b.xlsx", buf.getvalue())
    assert "Sales" in doc.tables
    assert "Sales" in doc.segments[0].location


def test_docx_sections_and_tables():
    import docx
    d = docx.Document()
    d.add_heading("Intro", 1)
    d.add_paragraph("Body text here")
    t = d.add_table(rows=2, cols=2)
    for i, v in enumerate(["h1", "h2", "v1", "v2"]):
        t.rows[i // 2].cells[i % 2].text = v
    buf = io.BytesIO()
    d.save(buf)
    doc = load_document("r.docx", buf.getvalue())
    assert any("Body text here" in s.text and s.location == "section: Intro" for s in doc.segments)
    assert "table 1" in doc.tables


def test_pptx_slides():
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Hello slide"
    buf = io.BytesIO()
    prs.save(buf)
    doc = load_document("s.pptx", buf.getvalue())
    assert doc.segments[0].location == "slide 1"
    assert "Hello slide" in doc.segments[0].text


def _pdf(text):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    if text:
        c.drawString(72, 720, text)
    c.showPage()
    c.save()
    return buf.getvalue()


def test_pdf_text_and_page_number():
    doc = load_document("p.pdf", _pdf("Page one text"))
    assert doc.segments[0].location == "page 1"
    assert "Page one text" in doc.segments[0].text


def test_blank_pdf_gives_clear_error():
    with pytest.raises(DocumentError):
        load_document("scan.pdf", _pdf(""))


def test_corrupt_file_raises_readable_error():
    with pytest.raises(DocumentError):
        load_document("bad.docx", b"PK-not-really")


def test_ragged_rows_to_df():
    # pyrefly: ignore [missing-import]
    from utils.document_loader import _rows_to_df
    # Header has 2 cols, row 1 has 3 cols, row 2 has 1 col
    rows = [["A", "B"], ["1", "2", "3"], ["4"]]
    df = _rows_to_df(rows)
    assert len(df.columns) == 3
    assert len(df) == 2
    assert df.iloc[0, 2] == "3"
    assert df.iloc[1, 1] == ""


def test_load_document_none_data():
    with pytest.raises(DocumentError):
        load_document("test.txt", None)