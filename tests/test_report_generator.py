import io

import pandas as pd
from pypdf import PdfReader

# pyrefly: ignore [missing-import]
from models import ReportData
# pyrefly: ignore [missing-import]
from utils.export_utils import available_formats, export_report
# pyrefly: ignore [missing-import]
from utils.report_generator import build_docx, build_markdown, build_pdf, build_text

LOGO = ('<svg xmlns="http://www.w3.org/2000/svg" width="64" height="64" viewBox="0 0 64 64">'
        '<rect x="10" y="4" width="38" height="52" rx="5" fill="#EEF2FF" stroke="#4F46E5" stroke-width="3"/>'
        '<circle cx="30" cy="30" r="6" fill="#4F46E5"/></svg>')


def _report(table=True):
    df = pd.DataFrame({"name": ["A", None], "value": [1.5, float("nan")]}) if table else None
    return ReportData(
        title="Quarterly Report \u00e9 <b> & more \u2713",
        summary="Short summary.",
        sections=[("Overview", "Intro paragraph with **bold**.\n\n- point one\n- point two\n\n| a | b |\n|---|---|\n| 1 | 2 |")],
        documents=[{"Document": "a.pdf", "Type": "PDF"}],
        table=df, sources=["a.pdf | page 1"], methodology="- Local parsing, Gemini analysis.")


def test_pdf_is_valid_and_contains_title_text():
    data = build_pdf(_report(), LOGO)
    assert data.startswith(b"%PDF")
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "Quarterly Report" in text and "Executive Summary" in text


def test_pdf_embeds_logo_graphics():
    with_logo = PdfReader(io.BytesIO(build_pdf(_report(), LOGO))).pages[0].get_contents().get_data()
    without = PdfReader(io.BytesIO(build_pdf(_report(), None))).pages[0].get_contents().get_data()
    assert len(with_logo) > len(without)


def test_pdf_handles_long_tables():
    report = _report()
    report.table = pd.DataFrame({"n": range(1000), "text": ["row " * 20] * 1000})
    assert len(PdfReader(io.BytesIO(build_pdf(report, LOGO))).pages) > 2


def test_docx_contains_headings():
    import docx
    d = docx.Document(io.BytesIO(build_docx(_report())))
    assert any("Executive Summary" in p.text for p in d.paragraphs)


def test_markdown_and_text_exports():
    assert "## 1. Executive Summary" in build_markdown(_report())
    assert "**" not in build_text(_report())


def test_formats_depend_on_content():
    assert "csv" in available_formats(_report(True))
    assert "csv" not in available_formats(_report(False))


def test_csv_neutralizes_formulas():
    report = _report()
    report.table = pd.DataFrame({"x": ["=SUM(A1:A2)"]})
    data, _, name = export_report("csv", report, None)
    assert b"'=SUM" in data and name.endswith(".csv")


def test_empty_tables_do_not_crash_rendering():
    # pyrefly: ignore [missing-import]
    from utils.report_generator import _rows_to_md, _docx_table
    import docx
    assert _rows_to_md([]) == ""
    d = docx.Document()
    _docx_table(d, [])  # Should not raise


def test_export_report_none_raises():
    import pytest
    with pytest.raises(ValueError):
        export_report("pdf", None, None)